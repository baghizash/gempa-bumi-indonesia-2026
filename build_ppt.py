#!/usr/bin/env python3
"""Build Portofolio-Analisis-Gempa-Bumi.pptx (16:9, 10 slides, ID).
Redesign T4 Brutalist Ringan Swiss: ground oat terang, headline Noto Sans
bold uppercase, blok warna penuh muted red / pumpkin, sudut tajam semua,
tanpa gradien / shadow / dekorasi. Angka dibaca dari findings.json dan
identik dengan deck lama.
"""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
F = json.loads((HERE / "findings.json").read_text())
CHARTS = HERE / "charts"
OUT = Path.home() / "workspace" / "your_files" / "Portofolio-Analisis-Gempa-Bumi.pptx"
OUT2 = HERE / "Portofolio-Analisis-Gempa-Bumi.pptx"

GROUND = RGBColor(0xF1, 0xE9, 0xDA)
SURFACE = RGBColor(0xF8, 0xF2, 0xE5)
INK = RGBColor(0x33, 0x23, 0x1C)
PRIMARY = RGBColor(0xA3, 0x4A, 0x38)
ACCENT = RGBColor(0xDE, 0x7A, 0x2D)
REDSOFT = RGBColor(0xC9, 0x7B, 0x63)
MUTE = RGBColor(0x8A, 0x7A, 0x6C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Noto Sans"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def fmt(n):
    return f"{n:,}".replace(",", ".")


def idnum(x, dec=1):
    return f"{x:.{dec}f}".replace(".", ",")


def bg(slide, color=GROUND):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, l, t, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def textbox(slide, l, t, w, h):
    tx = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.word_wrap = True
    return tf


def add_par(tf, text, size, color, bold=False, italic=False, first=False,
            align=None, space_after=8, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if line is not None:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return p


def add_bullet(tf, text, size=16, first=False):
    p = add_par(tf, text, size, INK, first=first, space_after=10, line=1.08)
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))
    pPr.set("marL", "171450")
    pPr.set("indent", "-171450")
    return p


def bullets(slide, items, l, t, w, h, size=16):
    tf = textbox(slide, l, t, w, h)
    for i, b in enumerate(items):
        add_bullet(tf, b, size=size, first=(i == 0))
    return tf


def title_block(slide, text, bar_color):
    tf = textbox(slide, 0.7, 0.42, 11.93, 0.85)
    add_par(tf, text.upper(), 31, INK, bold=True, first=True, space_after=0)
    rect(slide, 0.72, 1.3, 1.5, 0.09, bar_color)


def bottom_strip(slide):
    rect(slide, 0, 7.38, 13.333, 0.12, PRIMARY)


def caption(slide, text):
    tf = textbox(slide, 0.7, 6.88, 11.93, 0.4)
    add_par(tf, text, 11, MUTE, italic=True, first=True, space_after=0)


def chart_panel(slide, chart, l, t, w, h):
    rect(slide, l, t, w, h, SURFACE)
    path = CHARTS / f"{chart}.png"
    iw, ih = Image.open(path).size
    maxw, maxh = w - 0.3, h - 0.3
    scale = min(maxw / iw, maxh / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), Inches(l + (w - pw) / 2),
                             Inches(t + (h - ph) / 2), width=Inches(pw), height=Inches(ph))


def stat_block(slide, l, t, w, h, number, label, block_color, text_color, num_size=30):
    rect(slide, l, t, w, h, block_color)
    tf = textbox(slide, l + 0.25, t + 0.14, w - 0.5, h - 0.28)
    add_par(tf, number, num_size, text_color, bold=True, first=True, space_after=2)
    add_par(tf, label, 11, text_color, bold=True, space_after=0, line=1.05)


M = F["meta"]
MG = F["magnitudo"]
KD = F["kedalaman"]
WL = F["wilayah"]
TOP = F["top10_terbesar"]
BM = F["bmkg"]
DK = KD["Dangkal (<70 km)"]
WA = F["wilayah_teraktif"]
RATA = fmt(int(F["rata_per_tahun_penuh"]))

# ---- Slide 1: Cover, Asymmetric Split ----
s = prs.slides.add_slide(blank)
bg(s)
rect(s, 8.55, 0, 4.783, 7.5, PRIMARY)
tf = textbox(s, 8.95, 0.75, 4.0, 0.5)
add_par(tf, "GEMPA M≥4,5 TERCATAT", 13, GROUND, bold=True, first=True, space_after=0)
tf = textbox(s, 8.95, 1.2, 4.1, 1.5)
add_par(tf, fmt(M["total_kejadian"]), 68, WHITE, bold=True, first=True, space_after=0)
rect(s, 8.97, 2.85, 1.1, 0.08, ACCENT)
tf = textbox(s, 8.95, 3.25, 4.0, 1.1)
add_par(tf, "M" + idnum(MG["terbesar"]["magnitudo"]), 44, WHITE, bold=True, first=True, space_after=2)
add_par(tf, "GEMPA TERBESAR: ENDE, NTT (2026-08-15), KEDALAMAN 10 KM", 11.5, GROUND, bold=True, space_after=0, line=1.1)
rect(s, 8.55, 6.5, 4.783, 1.0, ACCENT)
tf = textbox(s, 8.95, 6.82, 4.0, 0.5)
add_par(tf, "SNAPSHOT BMKG: 20 OKTOBER 2026", 12.5, INK, bold=True, first=True, space_after=0)
rect(s, 0.7, 0.85, 0.28, 0.28, ACCENT)
tf = textbox(s, 0.7, 1.35, 7.5, 2.0)
add_par(tf, "GEMPA BUMI INDONESIA", 43, INK, bold=True, first=True, space_after=0)
add_par(tf, "2015-2026", 43, PRIMARY, bold=True, space_after=0)
tf = textbox(s, 0.7, 3.3, 7.3, 1.0)
add_par(tf, "Seberapa Sering, Seberapa Besar, dan Di Mana Gempa Mengguncang Indonesia",
        19, INK, first=True, space_after=0, line=1.15)
tf = textbox(s, 0.7, 4.5, 7.2, 1.3)
add_par(tf, f"Katalog USGS {fmt(M['total_kejadian'])} gempa M≥4,5, 1 Januari 2015 sampai "
            "20 Oktober 2026, plus snapshot real-time BMKG.", 14.5, INK, first=True, line=1.15)
add_par(tf, "Plus: paket dashboard Power BI (star schema, DAX, panduan rakit).", 14.5, INK, line=1.15)
rect(s, 0.7, 6.18, 1.6, 0.07, PRIMARY)
tf = textbox(s, 0.7, 6.38, 7.5, 0.55)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 17, INK, bold=True, first=True, space_after=0)

# ---- Slide 2: Latar belakang ----
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Kenapa Gempa Bumi?", PRIMARY)
bullets(s, [
    "Indonesia berdiri di pertemuan empat lempeng tektonik: Indo-Australia, "
    "Eurasia, Pasifik, dan Laut Filipina. Gempa bukan kejadian langka: "
    "ia terjadi hampir setiap hari.",
    "Tapi 'sering' saja tidak cukup untuk kebijakan dan kesiapsiagaan. Yang perlu "
    "dijawab: wilayah mana paling aktif, seberapa besar gempa yang realistis "
    "terjadi, dan seberapa dangkal sumbernya.",
    "Analisis ini menjawabnya dari katalog 12 tahun: frekuensi per tahun, kelas "
    "magnitudo, kelas kedalaman, peta sebaran, sampai gempa berflag tsunami, "
    "ditutup snapshot kejadian terbaru dari BMKG.",
], 0.7, 1.6, 7.5, 4.9)
stat_block(s, 8.6, 1.6, 4.0, 2.2, "4", "LEMPENG TEKTONIK BERTEMU DI INDONESIA", PRIMARY, WHITE, num_size=54)
stat_block(s, 8.6, 4.0, 4.0, 2.2, "12", "TAHUN KATALOG USGS: 2015-2026", ACCENT, INK, num_size=54)
caption(s, "Sumber data: USGS FDSN Event API (katalog historis) dan BMKG TEWS (snapshot real-time), diambil 20 Oktober 2026.")
bottom_strip(s)

# ---- Slide 3: Data dan metode ----
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Data dan Metode", ACCENT)
bullets(s, [
    f"Katalog mentah USGS: {fmt(M['katalog_mentah_usgs'])} kejadian dalam kotak batas Indonesia "
    f"(lintang -11 s.d. 6, bujur 95 s.d. 141), ambang M≥4,5. Setelah {fmt(M['dikeluarkan_label_negara_lain'])} kejadian "
    "berlabel negara lain (Filipina, Malaysia, Papua Nugini, Palau) dikeluarkan, katalog bersih berisi "
    f"{fmt(M['total_kejadian'])} kejadian.",
    f"{M['berlabel_timor_leste_dipertahankan']} kejadian berlabel 'Timor Leste' dipertahankan: episentrumnya di perairan "
    "Indonesia utara/timur Pulau Timor, hanya kota acuannya yang lintas negara.",
    "Wilayah diturunkan dari nama laut pada label USGS, lalu kotak koordinat pulau "
    "terdekat. Kelas kedalaman: dangkal <70 km, menengah 70-300 km, dalam >300 km.",
    "Snapshot BMKG (autogempa, gempaterkini, gempadirasakan) dipakai untuk bagian "
    "terkini: endpoint itu hanya menyimpan kejadian terbaru, bukan arsip.",
], 0.7, 1.6, 11.93, 3.8)
stat_block(s, 0.7, 5.55, 3.91, 1.3, fmt(M["katalog_mentah_usgs"]), "KEJADIAN MENTAH DALAM KOTAK BATAS INDONESIA", SURFACE, INK)
stat_block(s, 4.71, 5.55, 3.91, 1.3, fmt(M["dikeluarkan_label_negara_lain"]), "DIKELUARKAN: BERLABEL NEGARA LAIN", ACCENT, INK)
stat_block(s, 8.72, 5.55, 3.91, 1.3, fmt(M["total_kejadian"]), "KEJADIAN BERSIH DIANALISIS", PRIMARY, WHITE)
bottom_strip(s)

# ---- Slide 4: Tren per tahun ----
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Frekuensi: Hampir 800 Gempa per Tahun", PRIMARY)
bullets(s, [
    f"Rata-rata tahun penuh 2015-2025: {RATA} gempa M≥4,5 per tahun. "
    "Tidak ada tren naik/turun yang jelas: aktivitas berfluktuasi mengikuti "
    "rangkaian gempa besar dan gempa susulannya.",
    f"Tahun tersibuk: 2019 dengan {fmt(F['per_tahun']['2019'])} kejadian; tersepi: 2016 ({fmt(F['per_tahun']['2016'])}).",
    f"2026 baru berjalan sampai 20 Oktober tetapi sudah mencatat {fmt(F['per_tahun']['2026'])} kejadian, "
    "setara tahun penuh yang sibuk, didorong rangkaian gempa besar tahun ini.",
], 0.7, 1.6, 5.5, 3.4)
tf = textbox(s, 0.7, 5.2, 5.5, 1.5)
add_par(tf, RATA, 40, PRIMARY, bold=True, first=True, space_after=2)
add_par(tf, "RATA-RATA GEMPA M≥4,5 PER TAHUN (TAHUN PENUH 2015-2025)", 11.5, INK, bold=True, space_after=0, line=1.1)
chart_panel(s, "01_per_tahun", 6.55, 1.6, 6.08, 4.6)
caption(s, "Chart: jumlah gempa per tahun. 2026 (oranye) bersifat parsial.")
bottom_strip(s)

# ---- Slide 5: Magnitudo ----
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Magnitudo: Yang Besar Itu Langka, Tapi Nyata", ACCENT)
chart_panel(s, "04_kelas_magnitudo", 0.7, 1.6, 6.08, 4.6)
bullets(s, [
    f"Dari {fmt(M['total_kejadian'])} gempa: {fmt(MG['m5plus'])} mencapai M≥5,0, "
    f"{fmt(MG['m6plus'])} mencapai M≥6,0, dan hanya {MG['m7plus']} yang mencapai M≥7,0.",
    "Median magnitudo hanya 4,7: sebagian besar gempa Indonesia kecil. "
    "Risiko terkonsentrasi di ekor distribusi: belasan gempa M7+ per dekade.",
    "Pola ini konsisten dengan hukum Gutenberg-Richter: setiap kenaikan 1 skala "
    "magnitudo, frekuensi turun kira-kira 10 kali lipat.",
], 7.1, 1.6, 5.53, 3.2)
stat_block(s, 7.1, 5.0, 1.78, 1.3, fmt(MG["m5plus"]), "M≥5,0", PRIMARY, WHITE, num_size=22)
stat_block(s, 9.0, 5.0, 1.78, 1.3, fmt(MG["m6plus"]), "M≥6,0", ACCENT, INK, num_size=22)
stat_block(s, 10.9, 5.0, 1.73, 1.3, str(MG["m7plus"]), "M≥7,0", REDSOFT, INK, num_size=22)
caption(s, "Chart: kejadian per kelas magnitudo, 2015-2026.")
bottom_strip(s)

# ---- Slide 6: Terbesar ----
t0, t2 = TOP[0], TOP[2]
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Gempa Terbesar dalam 12 Tahun", PRIMARY)
stat_block(s, 0.7, 1.6, 5.5, 1.55, "M" + idnum(t0["magnitudo"]),
           f"{t0['lokasi_usgs'].upper()} ({t0['tanggal']}), KEDALAMAN {t0['kedalaman_km']:.0f} KM",
           PRIMARY, WHITE, num_size=48)
bullets(s, [
    f"Terbesar: M{idnum(t0['magnitudo'])}, {t0['lokasi_usgs']} ({t0['tanggal']}), "
    f"kedalaman hanya {t0['kedalaman_km']:.0f} km. Besar dan dangkal: kombinasi paling berbahaya.",
    f"Nama yang paling diingat publik: M{idnum(t2['magnitudo'])} Palu, Sulawesi ({t2['tanggal']}): "
    "gempa dangkal 20 km yang memicu tsunami dan likuefaksi.",
    "10 daftar teratas didominasi kawasan timur (NTT, Tanimbar, Halmahera, Laut Banda) "
    "dan Sulawesi, bukan Jawa yang penduduknya terpadat.",
], 0.7, 3.45, 5.5, 3.2, size=15.5)
chart_panel(s, "09_top10_terbesar", 6.55, 1.6, 6.08, 4.6)
caption(s, "Chart: 10 gempa terbesar 2015-2026 berdasarkan katalog USGS.")
bottom_strip(s)

# ---- Slide 7: Kedalaman ----
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Kedalaman: 7 dari 10 Gempa Itu Dangkal", ACCENT)
bullets(s, [
    f"Gempa dangkal (<70 km): {fmt(DK['jumlah'])} kejadian ({idnum(DK['persen'])}%). "
    f"Menengah: {fmt(KD['Menengah (70-300 km)']['jumlah'])} ({idnum(KD['Menengah (70-300 km)']['persen'])}%). "
    f"Dalam (>300 km): {fmt(KD['Dalam (>300 km)']['jumlah'])} ({idnum(KD['Dalam (>300 km)']['persen'])}%).",
    "Gempa dangkal lebih berpotensi merusak: energi dilepas dekat permukaan dan "
    "permukiman. Kedalaman median katalog hanya 36,2 km.",
    f"Kombinasi paling waspada: dangkal sekaligus M≥6, terjadi {KD['dangkal_m6plus']} kali "
    "dalam 12 tahun, kira-kira 9 kali per tahun.",
], 0.7, 1.6, 5.5, 3.4)
tf = textbox(s, 0.7, 5.2, 5.5, 1.5)
add_par(tf, idnum(DK["persen"]) + "%", 40, PRIMARY, bold=True, first=True, space_after=2)
add_par(tf, f"GEMPA DANGKAL (<70 KM): {fmt(DK['jumlah'])} DARI {fmt(M['total_kejadian'])} KEJADIAN",
        11.5, INK, bold=True, space_after=0, line=1.1)
chart_panel(s, "05_kelas_kedalaman", 6.55, 1.6, 6.08, 4.6)
caption(s, "Chart: jumlah kejadian per kelas kedalaman.")
bottom_strip(s)

# ---- Slide 8: Peta & wilayah ----
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Peta Sebaran dan Wilayah Teraktif", PRIMARY)
bullets(s, [
    f"Wilayah teraktif: {WA}, {fmt(WL[WA]['jumlah'])} kejadian, magnitudo maks M{idnum(WL[WA]['mag_maks'])}. "
    "Zona tumbukan Laut Maluku memang salah satu kawasan seismik paling aktif di dunia.",
    f"Menyusul: Sumatera ({fmt(WL['Sumatera']['jumlah'])}), Papua ({fmt(WL['Papua']['jumlah'])}), "
    f"Sulawesi ({fmt(WL['Sulawesi']['jumlah'])}). Kalimantan hanya {WL['Kalimantan']['jumlah']} kejadian: "
    "jauh dari batas lempeng, nyaris bebas gempa.",
    "Pola peta memperlihatkan jalur busur kepulauan: gempa menumpuk mengikuti "
    "garis subduksi dari barat Sumatera, selatan Jawa, sampai Banda dan Maluku.",
], 0.7, 1.6, 4.95, 5.0, size=15)
chart_panel(s, "08_peta_sebaran", 5.85, 1.6, 6.78, 4.6)
caption(s, "Chart: peta sebaran, ukuran titik = magnitudo, warna = kelas kedalaman.")
bottom_strip(s)

# ---- Slide 9: Snapshot BMKG ----
okt = BM["okt2026_dari_terkini"]
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Snapshot BMKG: Oktober 2026", ACCENT)
chart_panel(s, "12_snapshot_bmkg", 0.7, 1.6, 6.08, 3.3)
stat_block(s, 0.7, 5.1, 6.08, 1.3, "M6,1", "TERBESAR OKTOBER 2026: KODI, SUMBA BARAT DAYA (4 OKT)", PRIMARY, WHITE, num_size=30)
bullets(s, [
    f"Dalam 18 hari pertama Oktober 2026 saja, BMKG mencatat {len(okt)} gempa M5+, "
    "termasuk M6,1 di Kodi, Sumba Barat Daya (4 Okt) dan M5,8 di Tahuna, Kep. Sangihe (8 Okt).",
    "Dua gempa M5,7 berurutan: Sukabumi, Jawa Barat (11 Okt, dangkal 10 km) dan "
    "Pohuwato, Gorontalo (12 Okt). Terakhir di daftar: M5,3 Labuan Bajo, NTT (18 Okt).",
    f"Gempa terbaru saat data diambil: M{idnum(BM['gempa_terbaru_saat_ambil']['magnitudo'])} di perairan "
    "Pulau Saringi (20 Okt, 18:28 WIB), dirasakan III di Kabupaten Sumbawa. "
    "Bukti bahwa katalog ini hidup: gempa terjadi bahkan saat analisis dikerjakan.",
], 7.1, 1.6, 5.53, 4.9)
caption(s, "Chart: 15 gempa M5+ terbaru versi BMKG (snapshot 20 Oktober 2026).")
bottom_strip(s)

# ---- Slide 10: Penutup (Designed Close) ----
s = prs.slides.add_slide(blank)
bg(s)
title_block(s, "Penutup", PRIMARY)
stat_block(s, 0.7, 1.6, 3.91, 1.3, RATA, "GEMPA M≥4,5 PER TAHUN, 71,3% DI ANTARANYA DANGKAL", PRIMARY, WHITE, num_size=28)
stat_block(s, 4.71, 1.6, 3.91, 1.3, str(MG["m7plus"]), "GEMPA M≥7,0 DALAM 12 TAHUN: LANGKA, TANPA JADWAL", ACCENT, INK, num_size=28)
stat_block(s, 8.72, 1.6, 3.91, 1.3, fmt(WL[WA]["jumlah"]), "KEJADIAN DI HALMAHERA DAN MALUKU UTARA, WILAYAH TERAKTIF", REDSOFT, INK, num_size=28)
tf = textbox(s, 0.7, 3.3, 11.93, 1.0)
add_par(tf, "TERIMA KASIH", 44, INK, bold=True, first=True, space_after=0)
tf = textbox(s, 0.7, 4.35, 11.93, 1.1)
add_par(tf, "Risiko besar terkonsentrasi di timur: Halmahera-Maluku, Banda, NTT, Sulawesi. "
            "Paket dashboard Power BI tersedia di folder powerbi/ (star schema 3 tabel, measure DAX "
            "siap tempel dengan tabel verifikasi, layout 4 halaman termasuk halaman peta, dan panduan "
            "rakit langkah demi langkah). Semua angka dapat direproduksi: jalankan scripts/fetch_data.py "
            "lalu scripts/analyze.py, dan bisa dicek silang dengan sql/queries.sql.",
        13.5, INK, first=True, line=1.2, space_after=0)
rect(s, 0.7, 5.75, 1.6, 0.07, ACCENT)
tf = textbox(s, 0.7, 5.95, 11.93, 1.2)
add_par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 26, PRIMARY, bold=True, first=True, space_after=4)
add_par(tf, "github.com/baghizash/gempa-bumi-indonesia-2026", 15, INK, space_after=0)
bottom_strip(s)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
prs.save(str(OUT2))
print(f"PPT saved: {OUT}")
print(f"PPT saved: {OUT2}")
